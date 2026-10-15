"""
Actualiza el dashboard de encuestas post-llamada (NPS CLARO ECUADOR.html) con el Excel de NPS.

Uso:
    py actualizar_nps.py [ruta_excel] [--sin-subir]

Sin argumento toma el .xlsx más reciente de esta carpeta cuyo nombre empiece con "NPS"
(ej. "NPS 18.09 A 13.10 DIARIO FINAL.xlsx"). Solo reemplaza los datos (const D = {...});
el diseño y los cálculos del HTML no se tocan. Al terminar sube el HTML al repositorio
Reporte-NPS (git pull, commit y push con las credenciales guardadas en Windows), salvo
con --sin-subir.

- Se lee la hoja de detalle (la que tiene las columnas Tipo, Calificacion e ID): una fila por
  pregunta respondida. Se agrupa por ID de llamada: NPS, FCR, EPA (Amabilidad), CES -> CS y OC.
- Casos analizados: filas con Causa Raiz y Observacion. El motivo agrupado se conserva del
  HTML actual para las observaciones ya conocidas; las nuevas quedan como "Sin clasificar"
  y se listan al terminar para asignarles motivo.
"""
import collections
import datetime
import glob
import json
import os
import re
import subprocess
import sys

from openpyxl import load_workbook

CARPETA = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(CARPETA, "NPS CLARO ECUADOR.html")
DATOS_HTML = re.compile(r"const D = (\{.*?\});(\r?\n)", re.S)

# Tipo de pregunta (inicio del texto) -> posición de la métrica en el HTML (NPS, FCR, EPA, CS, OC)
METRICAS = {"NPS": 0, "FCR": 1, "EPA": 2, "CES": 3, "OC": 4}
CAUSAS = {"ASESOR": "Agente", "AGENTE": "Agente", "PROCESO": "Proceso", "CLIENTE": "Cliente",
          "PROCESO / ASESOR": "Proceso / Agente", "PROCESO / AGENTE": "Proceso / Agente"}
SIN_CLASIFICAR = "Sin clasificar"


def texto(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def clave(v):
    return texto(v).upper()


def buscar_excel():
    archivos = [f for f in glob.glob(os.path.join(CARPETA, "NPS*.xlsx"))
                if not os.path.basename(f).startswith("~$")]
    if not archivos:
        raise FileNotFoundError(f"No hay archivos NPS*.xlsx en {CARPETA}")
    return max(archivos, key=os.path.getmtime)


def hoja_detalle(wb):
    for ws in wb.worksheets:
        filas = ws.iter_rows(values_only=True)
        enc = [clave(h) for h in next(filas, [])]
        if {"TIPO", "CALIFICACION", "ID"} <= set(enc):
            return ws, enc
    raise RuntimeError("No encuentro la hoja de detalle (columnas Tipo, Calificacion e ID)")


def fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", texto(v)) or re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto(v))
    if not m:
        return None
    a, b, c = m.groups()
    return f"{a}-{b}-{c}" if len(a) == 4 else f"{c}-{int(b):02d}-{int(a):02d}"


def construir(ruta, anterior):
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws, enc = hoja_detalle(wb)
    col = {k: enc.index(k) for k in enc if k}

    def c(r, nombre, *alternativas):
        for n in (nombre,) + alternativas:
            if n in col and col[n] < len(r):
                return r[col[n]]
        return None

    llamadas, casos, descartados = {}, [], collections.Counter()
    filas = ws.iter_rows(values_only=True)
    next(filas)
    for r in filas:
        cid = texto(c(r, "ID"))
        if not cid:
            continue
        f = fecha(c(r, "FECHA COMPLETA")) or fecha(c(r, "FECHA"))
        if not f:
            descartados["sin fecha"] += 1
            continue
        # el nombre del asesor va tal cual (con sus espacios dobles): es la clave que usa el HTML
        ll = llamadas.setdefault(cid, {"orden": texto(c(r, "FECHA")), "fecha": f, "ase": str(c(r, "ASESOR") or "").strip(),
                                      "sup": texto(c(r, "SUPERVISOR")), "cola": texto(c(r, "COLA")),
                                      "cierre": texto(c(r, "CIERRE")), "m": [None] * 5})
        tipo = clave(c(r, "TIPO"))
        k = next((v for p, v in METRICAS.items() if tipo.startswith(p)), None)
        nota = c(r, "CALIFICACION")
        if k is not None and nota not in (None, ""):
            ll["m"][k] = int(float(nota))
        causa, obs = clave(c(r, "CAUSA RAIZ")), texto(c(r, "OBSERVACION"))
        if causa or obs:
            if k is None or causa not in CAUSAS or not obs:
                descartados["caso incompleto (causa/observación/métrica)"] += 1
            else:
                casos.append((cid, k, CAUSAS[causa], obs, ll["m"][k] if nota in (None, "") else int(float(nota))))
    wb.close()

    # motivos ya asignados en el HTML (por texto de la observación) y su texto original
    motivos = list(anterior.get("motivos", []))
    conocidos = {clave(a[4]): (a[3], a[4]) for a in anterior.get("an", [])}

    orden = sorted(llamadas.items(), key=lambda kv: (kv[1]["fecha"], kv[1]["orden"], kv[0]))
    asesores = sorted({v["ase"] for _, v in orden})
    sups = sorted({v["sup"] for _, v in orden if v["sup"]})
    colas = sorted({v["cola"] for _, v in orden})
    cierres = sorted({v["cierre"] for _, v in orden})
    fila = {}
    rows = []
    for i, (cid, v) in enumerate(orden):
        fila[cid] = i
        rows.append([v["fecha"], asesores.index(v["ase"]), sups.index(v["sup"]) if v["sup"] else None,
                     colas.index(v["cola"]), cierres.index(v["cierre"])] + v["m"])

    an, nuevas = [], collections.Counter()
    for cid, k, causa, obs, nota in casos:
        if clave(obs) in conocidos:
            m, obs = conocidos[clave(obs)]
        else:
            if SIN_CLASIFICAR not in motivos:
                motivos.append(SIN_CLASIFICAR)
            m = motivos.index(SIN_CLASIFICAR)
            nuevas[obs] += 1
        an.append([fila[cid], k, causa, m, obs, nota])
    an.sort(key=lambda a: (a[1], a[0]))

    D = dict(anterior)
    D.update({"asesores": asesores, "sups": sups, "colas": colas, "cierres": cierres,
              "rows": rows, "motivos": motivos, "an": an})
    return D, descartados, nuevas


def actualizar(ruta=None, html=HTML):
    ruta = str(ruta) if ruta else buscar_excel()
    with open(html, encoding="utf-8", newline="") as f:
        contenido = f.read()
    m = DATOS_HTML.search(contenido)
    if not m:
        raise RuntimeError(f"No está 'const D = {{...}};' en {html}")
    D, descartados, nuevas = construir(ruta, json.loads(m.group(1)))
    js = json.dumps(D, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    nuevo = contenido[:m.start()] + f"const D = {js};{m.group(2)}" + contenido[m.end():]
    tmp = html + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        f.write(nuevo)
    os.replace(tmp, html)

    fechas = [r[0] for r in D["rows"]]
    print(f"NPS actualizado con {os.path.basename(ruta)}")
    print(f"  {len(D['rows'])} encuestas del {min(fechas)} al {max(fechas)} · {len(D['asesores'])} asesores · "
          f"{len(D['sups'])} supervisores · {len(D['an'])} casos analizados")
    for k, n in descartados.items():
        print(f"  Descartados: {n} {k}")
    if nuevas:
        print(f"  Observaciones nuevas sin motivo ({len(nuevas)}), quedan como '{SIN_CLASIFICAR}':")
        for t, n in nuevas.most_common():
            print(f"    {n} x {t}")


def git(*args):
    r = subprocess.run(["git", *args], cwd=CARPETA, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} falló:\n{(r.stderr or r.stdout).strip()}")
    return r.stdout.strip()


def traer_cambios():
    """Trae lo último de GitHub (la descarga del OUT también sube la presentación a este repositorio)."""
    git("pull", "--ff-only", "origin", "main")


def subir(ruta):
    nombre = os.path.basename(HTML)
    git("add", nombre)
    if not git("status", "--porcelain", "--", nombre):
        print("GitHub: el NPS no cambió, no hay nada que subir")
        return
    # misma identidad que usan las descargas del IN y el OUT, para que funcione en cualquier PC
    git("-c", "user.name=1406-20", "-c", "user.email=1406-20@users.noreply.github.com",
        "commit", "-m", f"Actualización NPS con {os.path.basename(ruta)}", "--", nombre)
    git("push", "origin", "main")
    print("Subido a GitHub: https://a365-calidad-claro-ec.github.io/Reporte-NPS/ (tarda unos minutos en verse)")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--sin-subir"]
    con_subida = "--sin-subir" not in sys.argv
    try:
        ruta = args[0] if args else buscar_excel()
        if con_subida:
            traer_cambios()
        actualizar(ruta)
        if con_subida:
            subir(ruta)
    except Exception as e:
        print("ERROR:", e)
        sys.exit(1)
